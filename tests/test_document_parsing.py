import asyncio
from io import BytesIO
from uuid import UUID

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook
from PIL import Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.domains.documents.pipeline.chunking import chunk_document
from app.domains.documents.pipeline.document import (
    ChunkOptions,
    DocumentIR,
    DocumentSource,
    ParseOptions,
)
from app.domains.documents.pipeline.models import ExtractedElements, ExtractionResult
from app.domains.documents.pipeline.nodes.chunk import chunk
from app.domains.documents.pipeline.nodes.normalize import normalize
from app.domains.documents.pipeline.nodes.parse import parse
from app.domains.documents.pipeline.normalization import normalize_document
from app.domains.documents.pipeline.parsers import parse_document
from app.domains.documents.pipeline.parsers.errors import ParseError

DOC_ID = UUID("4c10a229-9b3a-40bc-8611-e2f787540f03")


def source(tmp_path, filename, data):
    path = tmp_path / "input.bin"
    path.write_bytes(data)
    return DocumentSource(document_id=DOC_ID, original_filename=filename, path=path)


def pdf_bytes(*, image_only=False, mixed=False):
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(600, 800))
    if image_only:
        pdf.drawInlineImage(Image.new("RGB", (10, 10), "black"), 50, 500, 100, 100)
    else:
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(40, 750, "Incident analysis")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(40, 720, "Requests became slow.")
        pdf.setFont("Courier", 10)
        pdf.drawString(40, 690, "def retry():")
        pdf.drawString(40, 678, "    return False")
        pdf.setFont("Helvetica", 10)
        for y in [580, 600, 620]:
            pdf.line(40, y, 300, y)
        for x in [40, 170, 300]:
            pdf.line(x, 580, x, 620)
        pdf.drawString(45, 606, "Symptom")
        pdf.drawString(175, 606, "Response")
        pdf.drawString(45, 586, "Latency")
        pdf.drawString(175, 586, "Measure")
    pdf.showPage()
    if mixed:
        pdf.drawInlineImage(Image.new("RGB", (10, 10), "black"), 50, 500, 100, 100)
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_markdown_preserves_sections_fenced_code_table_links_and_lines(tmp_path):
    text = (
        "# Incident\r\n\r\n[report](https://example.com/incident)\r\n\r\n"
        "```python\r\ndef retry():\r\n    return False\r\n```\r\n\r\n"
        "| Symptom | Choice |\r\n| --- | --- |\r\n| slow | measure |\r\n\r\n"
        "## Outcome\r\n\r\nNot yet confirmed.\r\n"
    )
    parsed = parse_document(source(tmp_path, "원래 이름.MD", text.encode()))
    assert parsed.original_filename == "원래 이름.MD"
    assert [b.block_type for b in parsed.blocks] == [
        "heading",
        "paragraph",
        "code",
        "table",
        "heading",
        "paragraph",
    ]
    code = parsed.blocks[2]
    assert "    return False\r\n" in code.text
    assert code.metadata["line_start"] == 5
    assert "https://example.com/incident" in parsed.blocks[1].text
    assert parsed.blocks[3].metadata["rows"] == [["Symptom", "Choice"], ["slow", "measure"]]
    assert parsed.blocks[-1].section_path == ["Incident", "Outcome"]
    for block in parsed.blocks:
        lines = text.splitlines(keepends=True)
        assert block.text == "".join(
            lines[block.metadata["line_start"] - 1 : block.metadata["line_end"]]
        )


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp949"])
def test_txt_korean_encoding_logs_and_sequence(tmp_path, encoding):
    text = "한글 문맥\n\nERROR request failed\n  at handler.py:3\n\n결과 미확인\n"
    parsed = parse_document(source(tmp_path, "incident.txt", text.encode(encoding)))
    assert [b.block_type for b in parsed.blocks] == ["paragraph", "log", "paragraph"]
    assert parsed.blocks[0].text == "한글 문맥\n"
    assert parsed.blocks[1].location == "lines 3-4"
    assert "  at handler.py:3" in parsed.blocks[1].text


def test_docx_keeps_body_order_heading_code_merged_table_and_nested_content(tmp_path):
    doc = Document()
    doc.add_heading("Incident", 1)
    doc.add_paragraph("Before the table")
    doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("def retry():\n    return False", style="Code")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged header"
    table.cell(1, 0).text = "Latency"
    table.cell(1, 1).text = "Measure"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested detail"
    nested.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Deep detail"
    doc.add_paragraph("After the table")
    output = BytesIO()
    doc.save(output)
    parsed = parse_document(source(tmp_path, "incident.docx", output.getvalue()))
    assert [b.block_type for b in parsed.blocks] == [
        "heading",
        "paragraph",
        "code",
        "table",
        "paragraph",
    ]
    assert parsed.blocks[-1].text == "After the table"
    assert parsed.blocks[2].text == "def retry():\n    return False"
    table_block = parsed.blocks[3]
    assert table_block.metadata["cells"][1]["origin"] == [1, 1]
    assert "Nested detail" in table_block.text
    assert "Deep detail" in table_block.text
    assert table_block.location == "body item 4"


def test_xlsx_keeps_arbitrary_columns_sheet_coordinates_merges_and_formulas(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "자유 양식"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Any heading"
    sheet["A2"], sheet["B2"] = "Latency", "=1+1"
    sheet["A4"] = "Independent section"
    workbook.create_sheet("Other")["C3"] = "Another experience"
    output = BytesIO()
    workbook.save(output)
    parsed = parse_document(source(tmp_path, "incident.xlsx", output.getvalue()))
    assert [b.location for b in parsed.blocks] == [
        "자유 양식!A1:B2",
        "자유 양식!A4:B4",
        "Other!A3:C3",
    ]
    block = parsed.blocks[0]
    assert block.metadata["formulas"] == {"B2": "=1+1"}
    assert block.metadata["merged_ranges"] == ["A1:B1"]
    assert block.metadata["row_numbers"] == [1, 2]
    assert block.metadata["rows"][1] == ["Latency", "=1+1"]
    assert parsed.metadata["formula_policy"] == "preserved_not_evaluated"


def test_pdf_keeps_page_order_code_and_grid_table(tmp_path):
    parsed = parse_document(source(tmp_path, "incident.pdf", pdf_bytes()))
    assert parsed.blocks[0].block_type == "heading"
    code = next(b for b in parsed.blocks if b.block_type == "code")
    assert "def retry():" in code.text and "return False" in code.text
    assert "\n    return False" in code.text
    table = next(b for b in parsed.blocks if b.block_type == "table")
    assert table.metadata["rows"] == [["Symptom", "Response"], ["Latency", "Measure"]]
    assert all(b.metadata["page"] == 1 for b in parsed.blocks)
    assert sum("Latency" in b.text for b in parsed.blocks) == 1


def test_image_pdf_fails_explicitly_and_mixed_pdf_exposes_unread_pages(tmp_path):
    with pytest.raises(ParseError) as exc:
        parse_document(source(tmp_path, "scan.pdf", pdf_bytes(image_only=True)))
    assert exc.value.code == "TEXT_UNAVAILABLE"
    parsed = parse_document(source(tmp_path, "mixed.pdf", pdf_bytes(mixed=True)))
    assert parsed.metadata["page_count"] == 2
    assert "TEXT_UNAVAILABLE:p.2" in parsed.warnings


@pytest.mark.parametrize(
    "filename,data,code",
    [
        ("a.txt", b"", "EMPTY_FILE"),
        ("a.txt", b" \n\t", "EMPTY_CONTENT"),
        ("a.txt", b"hello\x00world", "INVALID_TEXT"),
        ("a.txt", b"\xff", "UNSUPPORTED_ENCODING"),
        ("a.pdf", b"not a PDF", "INVALID_FILE"),
        ("a.pdf", b"%PDF-1.7\ngarbage", "INVALID_FILE"),
        ("a.docx", b"not a ZIP", "INVALID_FILE"),
        ("a.xlsx", b"not a ZIP", "INVALID_FILE"),
        ("a.bin", b"content", "UNSUPPORTED_FORMAT"),
    ],
)
def test_invalid_empty_or_mislabeled_files_fail_with_reason(tmp_path, filename, data, code):
    with pytest.raises(ParseError) as exc:
        parse_document(source(tmp_path, filename, data))
    assert exc.value.code == code


def test_parser_resource_limits(tmp_path):
    with pytest.raises(ParseError) as exc:
        parse_document(source(tmp_path, "a.txt", b"long text"), ParseOptions(max_bytes=2))
    assert exc.value.code == "FILE_TOO_LARGE"
    workbook = Workbook()
    workbook.active["Z100"] = "Sparse far-away cell"
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ParseError) as exc:
        parse_document(
            source(tmp_path, "a.xlsx", output.getvalue()), ParseOptions(max_sheet_cells=100)
        )
    assert exc.value.code == "SHEET_TOO_LARGE"


def test_expanded_office_size_is_checked_before_loading(tmp_path):
    workbook = Workbook()
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ParseError) as exc:
        parse_document(
            source(tmp_path, "a.xlsx", output.getvalue()), ParseOptions(max_expanded_bytes=10)
        )
    assert exc.value.code == "ARCHIVE_TOO_LARGE"


def test_txt_sections_and_multiline_traceback_are_not_split(tmp_path):
    raw = (
        "# Incident\n\nTraceback (most recent call last):\n  at worker.py\n\n"
        "ValueError: broken\n\n# Followup\n\nMeasure first.\n"
    )
    parsed = parse_document(source(tmp_path, "a.txt", raw.encode()))
    assert [b.block_type for b in parsed.blocks] == ["heading", "log", "heading", "paragraph"]
    assert "worker.py\n\nValueError: broken" in parsed.blocks[1].text
    assert parsed.blocks[-1].section_path == ["Followup"]


def test_pdf_korean_text_layer_is_extracted(tmp_path):
    output = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    pdf = canvas.Canvas(output)
    pdf.setFont("HYSMyeongJo-Medium", 12)
    pdf.drawString(40, 700, "한글 문서에서 요청 지연을 확인했다.")
    pdf.save()
    parsed = parse_document(source(tmp_path, "한글.pdf", output.getvalue()))
    assert "한글 문서에서 요청 지연을 확인했다." in "\n".join(b.text for b in parsed.blocks)


def test_normalization_is_repeatable_and_preserves_source_registry(tmp_path):
    input_source = source(tmp_path, "name.md", b"# Context\n\nAn observation.\n")
    parsed = parse_document(input_source)
    first = normalize_document(parsed)
    assert first == normalize_document(parsed)
    assert first == normalize_document(parse_document(input_source))
    assert DocumentIR.model_validate_json(first.model_dump_json()) == first
    for raw, block in zip(parsed.blocks, first.blocks, strict=True):
        ref = first.evidence_index[block.evidence_id]
        assert block.text == raw.text == ref.excerpt
        assert block.location == ref.location
        assert ref.document_id == DOC_ID and ref.block_id == block.block_id
        assert ref.start_char == 0 and ref.end_char == len(block.text)
    changed = parsed.model_copy(
        update={"document_id": UUID("3f0a6b2a-a149-475b-96b2-fca2ec7eb322")}
    )
    assert normalize_document(changed).blocks[0].block_id != first.blocks[0].block_id


def test_chunking_preserves_atomic_blocks_and_heading_context(tmp_path):
    raw = (
        b"# Incident\n\nIntro.\n\n```python\ndef retry():\n    return False\n```\n\n"
        b"| A | B |\n| - | - |\n| 1 | 2 |\n"
    )
    ir = normalize_document(parse_document(source(tmp_path, "incident.md", raw)))
    chunks = chunk_document(ir, ChunkOptions(target_chars=20))
    assert [b.block_id for c in chunks for b in c.blocks] == [b.block_id for b in ir.blocks]
    assert all(c.section_path == ["Incident"] for c in chunks)
    assert chunks[0].blocks[0].block_type == "heading"
    assert chunks[0].blocks[1].text == "Intro.\n"
    assert any(c.oversized for c in chunks)
    assert any(b.block_type == "table" for c in chunks for b in c.blocks)
    assert chunk_document(ir, ChunkOptions(target_chars=20)) == chunks


def test_upstream_node_retries_clear_stale_extractions_without_mutating_original(tmp_path):
    item = ExtractedElements.model_validate(
        {
            "cause": {"status": "hypothesis", "value": "Possible cause", "evidence_refs": ["e"]},
            "evidence_index": {
                "e": {
                    "document_id": str(DOC_ID),
                    "block_id": "b",
                    "location": "p.1",
                    "excerpt": "Possible cause",
                }
            },
        }
    )
    extraction = ExtractionResult(experiences=[item])
    state = {
        "sources": [source(tmp_path, "name.txt", b"Some observation")],
        "extraction_result": extraction,
    }
    state.update(asyncio.run(parse(state)))
    state.update(normalize(state))
    state.update(chunk(state))
    assert state["chunks"][0].blocks[0].evidence_id in state["documents"][0].evidence_index
    assert state["extraction_result"].experiences == []
    assert extraction.experiences[0].cause.status == "hypothesis"


def test_parse_node_uses_a_worker_thread(tmp_path, monkeypatch):
    import importlib
    import threading

    module = importlib.import_module("app.domains.documents.pipeline.nodes.parse")
    original = module.parse_document
    threads = []

    def observe(*args):
        threads.append(threading.get_ident())
        return original(*args)

    monkeypatch.setattr(module, "parse_document", observe)
    result = asyncio.run(module.parse({"sources": [source(tmp_path, "a.txt", b"Some text")]}))
    assert result["parsed_documents"]
    assert threads and threads[0] != threading.get_ident()
