from io import BytesIO

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.domains.documents.pipeline.document import ParsedBlock, ParseOptions
from app.domains.documents.pipeline.parsers.errors import ParseError


def _table_rows(table):
    rows = []
    anchors = {}
    cells = []
    for row_number, row in enumerate(table.rows, 1):
        values = []
        for column, cell in enumerate(row.cells, 1):
            origin = anchors.setdefault(cell._tc, [row_number, column])
            nested = [_table_rows(child) for child in cell.tables]
            text = cell.text
            values.append(text)
            cells.append(
                {
                    "row": row_number,
                    "column": column,
                    "origin": origin,
                    "text": text,
                    "nested_tables": nested,
                }
            )
        rows.append(values)
    return {"rows": rows, "cells": cells}


def _render_rows(rows):
    return "\n".join("\t".join(row) for row in rows)


def _render_docx_table(table):
    text = _render_rows(table["rows"])
    for cell in table["cells"]:
        if cell["origin"] != [cell["row"], cell["column"]]:
            continue
        for nested in cell["nested_tables"]:
            text += "\n" + _render_docx_table(nested)
    return text


def parse_docx(data: bytes):
    document = Document(BytesIO(data))
    blocks = []
    headings = []
    for position, item in enumerate(document.iter_inner_content(), 1):
        if isinstance(item, Table):
            metadata = _table_rows(item)
            text = _render_docx_table(metadata)
            kind = "table"
        else:
            text = item.text
            style = item.style.style_id if item.style is not None else ""
            metadata = {"style": style, "links": [link.url for link in item.hyperlinks]}
            kind = "paragraph"
            if style.startswith("Heading") and style[7:].isdigit():
                level = int(style[7:])
                headings = [entry for entry in headings if entry[0] < level]
                headings.append((level, text))
                kind = "heading"
            elif "code" in style.lower() or "preformat" in style.lower():
                kind = "code"
        if text.strip():
            blocks.append(
                ParsedBlock(
                    block_type=kind,
                    text=text,
                    location=f"body item {position}",
                    section_path=[entry[1] for entry in headings],
                    metadata={"body_item": position, **metadata},
                )
            )
    return blocks, {}, []


def parse_xlsx(data: bytes, options: ParseOptions):
    workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > options.max_sheet_cells:
                raise ParseError("SHEET_TOO_LARGE", f"Sheet exceeds cell limit: {sheet.title}")
            pending = []

            def flush():
                if not pending:
                    return
                first, last = pending[0][0].row, pending[-1][0].row
                cell_range = f"A{first}:{get_column_letter(sheet.max_column)}{last}"
                rows = [["" if c.value is None else str(c.value) for c in row] for row in pending]
                formulas = {
                    c.coordinate: c.value for row in pending for c in row if c.data_type == "f"
                }
                links = {
                    c.coordinate: c.hyperlink.target for row in pending for c in row if c.hyperlink
                }
                blocks.append(
                    ParsedBlock(
                        block_type="table",
                        text=_render_rows(rows),
                        location=f"{sheet.title}!{cell_range}",
                        section_path=[sheet.title],
                        metadata={
                            "sheet": sheet.title,
                            "range": cell_range,
                            "rows": rows,
                            "row_numbers": list(range(first, last + 1)),
                            "columns": [
                                get_column_letter(i) for i in range(1, sheet.max_column + 1)
                            ],
                            "formulas": formulas,
                            "links": links,
                            "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                        },
                    )
                )
                pending.clear()

            for row in sheet.iter_rows():
                if any(cell.value is not None for cell in row):
                    pending.append(row)
                else:
                    flush()
            flush()
    finally:
        workbook.close()
    return blocks, {"formula_policy": "preserved_not_evaluated"}, []
